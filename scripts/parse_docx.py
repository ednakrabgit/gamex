import docx
import json
import re

def parse_docx_files():
    doc_q = docx.Document('แนวข้อสอบปลายภาค_ป1_ภาคเรียน1_2569_50แผ่น_ฉบับปรับปรุง.docx')
    doc_a = docx.Document('เฉลย_แนวข้อสอบปลายภาค_ป1_50แผ่น_ฉบับปรับปรุง.docx')

    # Parse answer key
    answers_by_sheet = {}
    current_sheet_num = None
    for p in doc_a.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        # Check sheet header e.g. "หน้า 1 | คณิตศาสตร์ | จำนวนนับ 11–20" or "หน้า 14 | ..."
        m = re.match(r'^หน้า\s*(\d+)\s*\|\s*([^|]+)\s*\|\s*(.+)$', txt)
        if m:
            current_sheet_num = int(m.group(1))
            answers_by_sheet[current_sheet_num] = {
                'subject': m.group(2).strip(),
                'topic': m.group(3).strip(),
                'answers': []
            }
            continue
        if current_sheet_num is not None:
            ans_m = re.match(r'^(\d+)\.\s*(.+)$', txt)
            if ans_m:
                answers_by_sheet[current_sheet_num]['answers'].append({
                    'q_num': int(ans_m.group(1)),
                    'text': ans_m.group(2).strip()
                })

    print(f"Parsed {len(answers_by_sheet)} sheets from answer key.")

    # Parse questions
    sheets_q = []
    current_q_sheet = []
    for p in doc_q.paragraphs:
        txt = p.text.strip()
        if 'แนวข้อสอบปลายภาคเรียนที่' in txt:
            if current_q_sheet:
                sheets_q.append(current_q_sheet)
                current_q_sheet = []
        if txt:
            current_q_sheet.append(txt)
    if current_q_sheet:
        sheets_q.append(current_q_sheet)

    print(f"Parsed {len(sheets_q)} question sheets from question doc.")

    parsed_sheets = []
    for idx, lines in enumerate(sheets_q):
        sheet_num = idx + 1
        subject = "ทั่วไป"
        topic = f"ชุดที่ {sheet_num}"
        instructions = ""
        questions = []

        for line in lines:
            if '|' in line and not line.startswith('1.') and not line.startswith('2.'):
                parts = line.split('|')
                if len(parts) >= 2:
                    subject = parts[0].strip()
                    topic = parts[1].strip()
            elif 'คำชี้แจง' in line:
                instructions = line.replace('คำชี้แจง:', '').strip()
            else:
                qm = re.match(r'^(\d+)\.\s*(.+)$', line)
                if qm:
                    q_num = int(qm.group(1))
                    q_text = qm.group(2).strip()
                    questions.append({
                        'number': q_num,
                        'question': q_text,
                        'raw_line': line
                    })

        # match with answer sheet
        sheet_ans_info = answers_by_sheet.get(sheet_num, {'answers': []})
        ans_list = sheet_ans_info.get('answers', [])
        ans_dict = {a['q_num']: a['text'] for a in ans_list}

        parsed_sheets.append({
            'sheet_num': sheet_num,
            'subject': subject,
            'topic': topic,
            'instructions': instructions,
            'questions': questions,
            'answers': ans_dict
        })

    return parsed_sheets

if __name__ == '__main__':
    sheets = parse_docx_files()
    print(f"Sample Sheet 1: {sheets[0]}")
    print(f"Sample Sheet 3: {sheets[2]}")
    with open('parsed_sheets_raw.json', 'w', encoding='utf-8') as f:
        json.dump(sheets, f, ensure_ascii=False, indent=2)
